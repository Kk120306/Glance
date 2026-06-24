#!/usr/bin/env python3
"""Generate Glance pitch deck as PowerPoint."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SLIDES = [
    {
        "title": "Glance",
        "subtitle": "Hands-free communication for people who cannot use a keyboard, mouse, or touchscreen.",
        "bullets": [
            "Real-time messaging between family and motor-impaired patients",
            "Reply with gaze and blink — never touch the screen",
            "Built for locked-in syndrome, ALS, cerebral palsy, and similar conditions",
        ],
        "is_title": True,
    },
    {
        "title": "The Problem",
        "bullets": [
            "Millions live with severe motor impairment and cannot type, tap, or point",
            "Families want to stay connected — but have no real-time channel that works without hands",
            "Commercial AAC devices cost $5,000–$20,000+ and require clinical setup",
            "Existing tools don't run on a laptop, webcam, and microphone you already own",
        ],
    },
    {
        "title": "What's at Stake",
        "bullets": [
            "Can't answer a simple yes/no question from a loved one in the moment",
            "Can't request water, pain relief, or help without a caregiver in the room",
            "Can't trigger emergency help when the camera is off",
            "Communication becomes slow, one-directional, and emotionally distant",
        ],
    },
    {
        "title": "Our Solution",
        "bullets": [
            "Glance is a hands-free AAC platform: receive messages, reply with eyes, signal distress anytime",
            "Patient app: animated guide reads messages aloud; gaze selects curated replies",
            "Family dashboard: compose messages, clone your voice, manage multiple patients",
            "Zero hands. Zero keyboard. Zero mouse. Ever.",
        ],
    },
    {
        "title": "How It Works",
        "bullets": [
            "Family sends a message — tone classified, delivered in their cloned voice",
            "Patient hears it via an expressive on-screen guide synced to speech",
            "Patient selects from AI-ranked phrase suggestions — always confirms before send",
            "Yes/no questions auto-detect: giant UP=YES / DOWN=NO targets",
            "SOS works without the camera via calibrated vocalization detection",
        ],
    },
    {
        "title": "Why Glance Wins",
        "bullets": [
            "Off-the-shelf hardware — browser, webcam, mic; no proprietary device",
            "Voice cloning in-app — messages feel personal, not robotic",
            "Ethical AI gate — curated suggestions only; patient always confirms",
            "ScanMode fallback when camera unavailable; SOS always works via microphone",
            "Privacy-aware camera windows with enforced track cleanup",
        ],
    },
    {
        "title": "Market Opportunity",
        "bullets": [
            "Global AAC market projected at $4B+ and growing with aging populations",
            "Beachhead: home-based care for ALS, locked-in syndrome, severe CP, stroke recovery",
            "Expand: skilled nursing facilities, rehab centers, telehealth care teams",
        ],
    },
    {
        "title": "Business Model",
        "bullets": [
            "SaaS subscription per patient household — dashboard + patient app + real-time alerts",
            "Tiered plans: single patient, multi-patient family, facility license",
            "Usage-based pass-through for voice synthesis at scale",
            "Partnerships with ALS associations, rehab networks, and payers",
        ],
    },
    {
        "title": "Traction & Status",
        "bullets": [
            "Feature-complete MVP: messaging, gaze tracking, voice clone, SOS, multi-patient dashboard",
            "213 automated tests; hard-constraint suite enforces zero-hands and SOS independence",
            "Phases 1–6 shipped: read receipts, offline queue, sender personas, media upload",
            "Next: real-device validation with pilot patients and caregivers",
        ],
    },
    {
        "title": "The Ask",
        "bullets": [
            "Pilot partners: 3–5 patient–caregiver pairs for real-world validation",
            "Advisors with AAC clinical experience and assistive-tech go-to-market",
            "Seed capital for pilots, HIPAA-ready infrastructure, and device testing lab",
            "Glance gives voice back to people the world stopped listening to.",
        ],
    },
]

ACCENT = RGBColor(0x25, 0x63, 0xEB)
DARK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_slide_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_footer(slide, text: str, prs: Presentation) -> None:
    left = Inches(0.6)
    top = prs.slide_height - Inches(0.45)
    width = prs.slide_width - Inches(1.2)
    box = slide.shapes.add_textbox(left, top, width, Inches(0.3))
    tf = box.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.font.size = Pt(10)
    p.font.color.rgb = MUTED
    p.alignment = PP_ALIGN.RIGHT


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for i, data in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(blank)
        set_slide_bg(slide, WHITE)

        # Accent bar
        bar = slide.shapes.add_shape(
            1,  # rectangle
            Inches(0),
            Inches(0),
            Inches(0.12),
            prs.slide_height,
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background()

        title_top = Inches(1.0) if not data.get("is_title") else Inches(2.2)
        title_box = slide.shapes.add_textbox(Inches(0.9), title_top, Inches(11.5), Inches(1.2))
        tf = title_box.text_frame
        tf.text = data["title"]
        p = tf.paragraphs[0]
        p.font.size = Pt(44 if data.get("is_title") else 36)
        p.font.bold = True
        p.font.color.rgb = DARK

        content_top = title_top + Inches(1.1)
        if data.get("subtitle"):
            sub_box = slide.shapes.add_textbox(Inches(0.9), content_top, Inches(11.0), Inches(0.9))
            stf = sub_box.text_frame
            stf.text = data["subtitle"]
            sp = stf.paragraphs[0]
            sp.font.size = Pt(20)
            sp.font.color.rgb = MUTED
            content_top += Inches(1.0)

        bullet_top = content_top + (Inches(0.3) if data.get("is_title") else Inches(0))
        bullet_box = slide.shapes.add_textbox(Inches(0.9), bullet_top, Inches(11.0), Inches(4.5))
        btf = bullet_box.text_frame
        btf.word_wrap = True
        btf.vertical_anchor = MSO_ANCHOR.TOP

        for j, bullet in enumerate(data["bullets"]):
            para = btf.paragraphs[0] if j == 0 else btf.add_paragraph()
            para.text = bullet
            para.level = 0
            para.font.size = Pt(22 if data.get("is_title") else 20)
            para.font.color.rgb = DARK
            para.space_after = Pt(14)
            para.line_spacing = 1.15

        add_footer(slide, f"Glance · Slide {i} of {len(SLIDES)}", prs)

    out = Path(__file__).resolve().parent / "Glance-Pitch-Deck.pptx"
    prs.save(out)
    return out


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
